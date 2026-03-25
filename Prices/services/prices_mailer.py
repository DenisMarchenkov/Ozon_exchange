import os
from typing import List
from pathlib import Path
from datetime import datetime
import pandas as pd

from Common.base_mailer import BaseMailer
from Common.settings import SUPPLIER_PRICE_FOLDER


class OzonPriceSyncMailer(BaseMailer):
    def __init__(
        self,
        missing_in_ozon: set,
        missing_in_price: set,
        filtered_prices: List[dict],
        *args,
        **kwargs
    ):
        super().__init__(*args, **kwargs)
        self.missing_in_ozon = missing_in_ozon
        self.missing_in_price = missing_in_price
        self.filtered_prices = filtered_prices

    # ---------- Subject ----------

    def build_subject_core(self) -> str:
        return "Отчет по обмену цен Ozon"

    # ---------- Plain text ----------

    def build_body(self) -> str:
        return (
            "Обмен цен с Ozon завершен.\n\n"
            f"Всего товаров отправлено: {len(self.filtered_prices)}\n"
            f"Есть в Ozon, но нет в прайсе: {len(self.missing_in_price)}\n"
            f"Есть в прайсе, но нет в Ozon: {len(self.missing_in_ozon)}\n\n"
            "Подробности во вложении."
        )

    # ---------- HTML ----------

    def build_body_html(self) -> str:
        return f"""
        <h2>Обмен цен с Ozon завершен</h2>

        <table style="border-collapse: collapse;">
            <tr>
                <td style="padding:8px; border:1px solid #ccc;">Отправлено товаров</td>
                <td style="padding:8px; border:1px solid #ccc;"><b>{len(self.filtered_prices)}</b></td>
            </tr>
            <tr>
                <td style="padding:8px; border:1px solid #ccc;">Есть в Ozon, но нет в прайсе</td>
                <td style="padding:8px; border:1px solid #ccc;"><b>{len(self.missing_in_price)}</b></td>
            </tr>
            <tr>
                <td style="padding:8px; border:1px solid #ccc;">Есть в прайсе, но нет в Ozon</td>
                <td style="padding:8px; border:1px solid #ccc;"><b>{len(self.missing_in_ozon)}</b></td>
            </tr>
        </table>

        <p>Подробная информация во вложенном файле.</p>

        {self.build_signature_html()}
        """

    # ---------- Excel генерация ----------

    def _build_excel_report(self) -> Path:
        from openpyxl import load_workbook
        from openpyxl.styles import PatternFill, Font, Alignment
        from openpyxl.utils import get_column_letter

        filename = f"ozon_price_sync_{datetime.now():%Y-%m-%d_%H-%M-%S}.xlsx"
        #file_path = Path(filename)
        file_path = os.path.join(SUPPLIER_PRICE_FOLDER, filename)

        # объединяем все offer_id
        all_offer_ids = self.missing_in_price | self.missing_in_ozon

        data = []
        for offer_id in all_offer_ids:
            in_price = offer_id in self.missing_in_price
            in_ozon = offer_id in self.missing_in_ozon

            # статус (человеческий)
            if in_price and not in_ozon:
                status = "Нет в прайсе"
            elif in_ozon and not in_price:
                status = "Нет в Ozon"
            elif in_price and in_ozon:
                status = "Конфликт"
            else:
                status = "ОК"

            data.append({
                'offer_id': offer_id,
                'missing_in_price': 'Да' if in_price else 'Нет',
                'missing_in_ozon': 'Да' if in_ozon else 'Нет',
                'status': status
            })

        # сортировка
        data = sorted(data, key=lambda x: x['offer_id'])

        # записываем через pandas
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            df = pd.DataFrame(data)
            df.to_excel(writer, sheet_name='differences', index=False)

            pd.DataFrame(self.filtered_prices).to_excel(
                writer,
                sheet_name='updated_prices',
                index=False
            )

        # --------- КРАСОТА через openpyxl ---------
        wb = load_workbook(file_path)
        ws = wb['differences']

        # стили
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

        bold_font = Font(bold=True)
        center_align = Alignment(horizontal="center")

        # заголовки
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = bold_font
            cell.alignment = center_align

        # автоширина колонок
        for col in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col[0].column)

            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))

            ws.column_dimensions[col_letter].width = max_length + 3

        # раскраска
        for row in ws.iter_rows(min_row=2):
            missing_price_cell = row[1]  # B
            missing_ozon_cell = row[2]  # C
            status_cell = row[3]  # D

            # Да / Нет цвета
            if missing_price_cell.value == 'Да':
                missing_price_cell.fill = red_fill
            else:
                missing_price_cell.fill = green_fill

            if missing_ozon_cell.value == 'Да':
                missing_ozon_cell.fill = red_fill
            else:
                missing_ozon_cell.fill = green_fill

            # статус
            if status_cell.value == "Конфликт":
                status_cell.fill = yellow_fill
            elif status_cell.value == "ОК":
                status_cell.fill = green_fill
            else:
                status_cell.fill = red_fill

        wb.save(file_path)

        return file_path

    # ---------- Attachments ----------

    def build_attachments(self) -> List[Path]:
        file_path = self._build_excel_report()
        return [file_path]