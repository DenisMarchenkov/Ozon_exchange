import os
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.worksheet.page import PageMargins


class WarehouseExcelFormatter:
    def __init__(self, file_path: str):
        self.file_name = os.path.basename(file_path)

    # ----------------------------------------------------
    # Общие свойства листа
    # ----------------------------------------------------
    def apply_common(self, sheet: Worksheet, suffix: str):
        cm = 1 / 2.54
        sheet.page_margins = PageMargins(
            left=cm * 0.8,
            right=cm * 0.8,
            top=cm * 0.8,
            bottom=cm * 1.8,
        )

        sheet.oddFooter.left.text = (
            f"{os.path.splitext(self.file_name)[0]} - {suffix}"
        )

        alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = alignment
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=14)

    def format_orders_summary(self, sheet: Worksheet):
        sheet.insert_rows(1)
        sheet["A1"].value = "СВОДКА ПО ЗАКАЗАМ"
        sheet["A1"].font = Font(name="Calibri", size=20, bold=True)

        sheet.column_dimensions["A"].width = 31
        sheet.column_dimensions["B"].width = 18
        sheet.column_dimensions["C"].width = 18
        sheet.column_dimensions["D"].width = 15
        sheet.column_dimensions["E"].width = 15

        sheet.print_title_rows = "2:2"

        date_format = "DD.MM.YYYY"

        for row in range(3, sheet.max_row + 1):
            sheet.cell(row, 4).number_format = date_format
            sheet.cell(row, 5).number_format = date_format

    def format_items_summary(self, sheet: Worksheet):
        sheet.insert_rows(1)
        sheet["A1"].value = "СВОДКА ПО БРЕНДАМ"
        sheet["A1"].font = Font(name="Calibri", size=20, bold=True)

        sheet.column_dimensions["A"].width = 15  # бренд
        sheet.column_dimensions["B"].width = 15  # артикул
        sheet.column_dimensions["C"].width = 45  # наименование
        sheet.column_dimensions["D"].width = 15  # срок годности
        sheet.column_dimensions["E"].width = 7  # количество

        sheet.print_title_rows = "2:2"

        grey = PatternFill("solid", fgColor="808080")
        white_font = Font(color="FFFFFF", size=14)
        thin_white = Border(
            bottom=Side(style="thin", color="FFFFFF"),
            right=Side(style="thin", color="FFFFFF"),
        )

        date_format = "DD.MM.YYYY"

        for row in range(3, sheet.max_row + 1):
            sheet.cell(row, 4).number_format = date_format

            # подсветка одинаковых артикулов
            cur = sheet.cell(row, 2).value
            next_val = (
                sheet.cell(row + 1, 2).value
                if row < sheet.max_row
                else None
            )

            if cur == next_val:
                for col in (2, 3, 4):
                    for r in (row, row + 1):
                        c = sheet.cell(r, col)
                        c.fill = grey
                        c.font = white_font
                        c.border = thin_white

    def format_full_data(self, sheet: Worksheet):
        sheet.page_setup.orientation = 'landscape'
        sheet.insert_rows(1)
        sheet["A1"].value = "ЛИСТ ПОДБОРА ЗАКАЗОВ"
        sheet["A1"].font = Font(name="Calibri", size=20, bold=True)

        sheet.column_dimensions["A"].width = 24  # номер заказа
        sheet.column_dimensions["B"].width = 15  # бренд
        sheet.column_dimensions["C"].width = 15  # артикул
        sheet.column_dimensions["D"].width = 40  # наименование
        sheet.column_dimensions["E"].width = 6  # количество
        sheet.column_dimensions["F"].width = 14 # цена с ндс
        sheet.column_dimensions["G"].width = 15 # срок годности
        sheet.column_dimensions["H"].width = 15 # дата заказа
        sheet.column_dimensions["I"].width = 15 # дата отгрузки

        sheet.print_title_rows = "2:2"
